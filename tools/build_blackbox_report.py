"""Build the report from successful, dated test evidence; no operational data."""
import json
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
evidence = ET.parse(ROOT / '.local/blackbox-20260917-phpunit.xml')
cases = {c.attrib['name']: c for c in evidence.findall('.//testcase')}
assert len(cases) == 46
assert sum(int(c.attrib.get('assertions', '0')) for c in cases.values()) == 401
assert not evidence.findall('.//failure') and not evidence.findall('.//error')
assert not evidence.findall('.//skipped')
browser = json.loads((ROOT / '.local/blackbox-20260917-browser.json').read_text(encoding='utf-8-sig'))
assert len(browser['passed']) == 8 and not browser['errors'] and not browser['failedLocalRequests']

keys = {
 'login': 'valid_login_and_logout', 'wrong': 'wrong_login_and_empty_credentials_are_rejected',
 'guest': 'guest_cannot_access_administration', 'deleted': 'deleted_account_cannot_login',
 'dashboard': 'dashboard_displays_correct_month_year_and_combined_counts',
 'pages': 'filters_lookup_and_all_pages_render',
 'crud': 'both_letter_types_complete_create_read_update_delete_through_http',
 'required': 'required_letter_fields_are_rejected_for_both_types',
 'limits': 'dates_and_text_outside_mysql_limits_return_validation_errors',
 'attachment': 'attachment_type_and_ten_megabyte_limit_are_validated',
 'search': 'search_month_year_all_months_and_date_order_for_both_types',
 'pagination': 'pagination_shows_ten_records_and_remaining_records_on_page_two',
 'reply': 'reply_links_are_optional_validated_and_visible_both_ways',
 'protected': 'referenced_incoming_letter_and_reset_are_protected',
 'reset': 'reset_requires_confirmation_and_keeps_other_letter_type',
 'csv': 'csv_preview_confirmation_and_replay_protection',
 'badimport': 'invalid_import_file_and_other_user_confirmation_are_rejected',
 'sheets': 'same_workbook_imports_incoming_sheet1_and_outgoing_sheet2',
 'merged': 'sheet2_legacy_agenda_handles_multiline_headers_and_merged_dates',
 'fallback': 'legacy_blank_outgoing_dates_use_letter_dates_and_skip_three_row_header',
 'invaliddate': 'explicit_invalid_legacy_outgoing_date_is_not_replaced',
 'bothdates': 'legacy_missing_both_dates_is_not_accepted',
 'flatdate': 'a_truly_missing_outgoing_date_still_requires_validation',
 'sheetname': 'sheet_name_does_not_override_wrong_letter_headers',
 'named': 'matching_named_sheet_is_preferred_case_insensitively',
 'nosheet': 'missing_outgoing_sheet_is_reported_without_reading_incoming_rows',
 'export': 'both_letter_exports_are_downloadable_excel_files',
 'profile': 'profile_updates_and_admin_creation',
 'email': 'duplicate_and_invalid_profile_admin_emails_are_rejected',
 'password': 'wrong_old_password_short_password_and_mismatched_confirmation_are_rejected',
 'account': 'account_removal_preserves_authorship_and_revokes_sessions',
 'forge': 'client_cannot_forge_creator_or_editor',
}

# Each line: feature/scenario | input/action | expected | observed | evidence keys.
data = """
Login dengan akun aktif|Email dan kata sandi benar.|Pengguna masuk ke dashboard.|Autentikasi berhasil dan dashboard dapat diakses.|login,b0
Login dengan kata sandi salah|Email terdaftar dan kata sandi salah.|Login ditolak dan pesan kesalahan ditampilkan.|Pesan validasi muncul; pengguna tetap belum login.|wrong
Validasi formulir login|Kredensial kosong atau format email salah.|Kolom wajib dan format email divalidasi.|Input ditolak dengan pesan kesalahan pada kolom terkait.|wrong
Pembatasan akses tamu|Buka halaman administrasi tanpa login.|Pengguna diarahkan ke halaman login.|Akses administrasi dialihkan ke login.|guest
Logout|Keluar dari akun yang sedang login.|Sesi autentikasi berakhir.|Pengguna keluar dan kembali ke halaman login.|login
Login akun nonaktif|Gunakan kredensial akun yang telah dihapus.|Akun tidak dapat login.|Autentikasi akun nonaktif ditolak.|deleted
Perhitungan dashboard|Fixture September 2026: 2 surat masuk bulan ini, 1 surat keluar bulan ini, 3 surat tahun ini, 5 surat keseluruhan.|Empat kartu menampilkan 2, 1, 3, dan 5.|Keempat angka sesuai; daftar surat terbaru tampil.|dashboard
Navigasi menu utama|Buka dashboard, surat masuk, surat keluar, profil, dan halaman impor.|Halaman tersedia dan dapat ditampilkan.|Seluruh halaman yang diperiksa berhasil dimuat.|pages,b1
Tambah surat masuk|Isi data surat masuk dengan tanggal dan perihal valid.|Surat tersimpan dan tampil pada daftar.|Surat baru tampil pada daftar dan memiliki halaman detail.|crud
Detail surat masuk|Buka detail surat masuk yang baru dibuat.|Data surat dan pencatat ditampilkan.|Isi surat serta nama pencatat tampil sesuai data uji.|crud,b3
Ubah surat masuk|Ubah perihal surat masuk yang tersimpan.|Perubahan ditampilkan setelah disimpan.|Detail menampilkan perihal yang telah diubah.|crud
Hapus surat masuk|Hapus surat masuk yang tidak memiliki balasan.|Surat hilang dari daftar dan detail tidak tersedia.|Surat tidak tampil; permintaan detail menghasilkan 404.|crud
Tambah surat keluar|Isi surat keluar biasa dengan data valid.|Surat tersimpan dan tampil pada daftar.|Surat baru tampil pada daftar dan memiliki detail.|crud
Detail surat keluar|Buka detail surat keluar yang baru dibuat.|Data surat dan pencatat ditampilkan.|Data surat dan nama pencatat tampil sesuai input.|crud
Ubah surat keluar|Ubah perihal surat keluar yang tersimpan.|Perubahan tersimpan dan ditampilkan.|Detail menampilkan perihal yang telah diubah.|crud
Hapus surat keluar|Hapus surat keluar yang tersimpan.|Surat hilang dari daftar dan detail tidak tersedia.|Surat tidak tampil; detail menghasilkan 404.|crud
Validasi kolom wajib surat|Kosongkan perihal, tanggal surat, serta tanggal masuk atau keluar.|Penyimpanan ditolak dengan pesan kolom wajib.|Tiga kolom tersebut menghasilkan kesalahan validasi pada kedua jenis surat.|required
Validasi batas tanggal dan teks|Tanggal 0999-01-01, tanggal keluar dengan waktu, atau perihal 16.001 karakter.|Input di luar aturan penyimpanan ditolak.|Kesalahan validasi diterima; input tidak disimpan.|limits
Unggah dan buka lampiran|Unggah gambar PNG pada surat masuk lalu buka lampiran.|Lampiran tersimpan dan dapat dibuka.|Pencatatan berhasil; lampiran memberikan respons HTTP 200.|b3
Penolakan jenis lampiran|Unggah berkas EXE pada kedua jenis surat.|Jenis berkas yang tidak diizinkan ditolak.|Pesan validasi lampiran muncul pada kedua formulir.|attachment
Batas ukuran lampiran|Unggah PDF berukuran 10.241 KB, melebihi batas 10 MB.|Berkas terlalu besar ditolak.|Pesan validasi ukuran lampiran muncul.|attachment
Pembacaan OCR|Jalankan OCR pada gambar uji yang memuat tulisan SURAT.|Tulisan pada gambar dapat dikenali.|OCR menghasilkan teks yang memuat SURAT.|b7
Pencarian surat|Cari penanda BB-FEB pada kedua daftar surat.|Hanya data yang cocok ditampilkan.|Surat Februari tampil; data Januari dan Maret tidak tampil.|search
Pencarian tanpa hasil|Cari teks yang tidak ada pada data uji.|Halaman menampilkan daftar kosong tanpa gagal.|Respons HTTP 200 dengan hasil kosong.|search
Filter bulan dan tahun|Pilih Februari dan tahun 2026 pada kedua daftar.|Hanya surat Februari 2026 ditampilkan.|Data Februari tampil; data bulan lain tidak tampil.|search
Urutan semua bulan|Tampilkan semua bulan untuk data Januari, Februari, dan Maret 2026.|Surat diurutkan berdasarkan tanggal surat masuk atau keluar dari terbaru.|Urutan Maret, Februari, Januari sesuai pada kedua daftar.|search
Validasi parameter filter|Gunakan bulan 99 atau tahun dua digit 20.|Parameter tidak valid ditolak.|Permintaan menghasilkan validasi parameter filter.|pages,search
Pembagian halaman|Tampilkan 12 surat lalu buka halaman kedua.|Halaman pertama berisi 10 surat; halaman kedua 2 surat.|Halaman pertama memuat surat 12–03; halaman kedua 02–01.|pagination
Hubungan surat balasan|Pilih surat masuk sebagai asal balasan surat keluar.|Referensi tersimpan; detail memiliki tautan dua arah.|Pencarian, pemilihan, serta tautan detail surat asal dan balasan berfungsi.|reply,b4
Referensi balasan opsional|Buat surat keluar biasa; hapus referensi balasan saat mengubah surat.|Surat keluar dapat disimpan tanpa surat asal.|Pembuatan biasa dan pengosongan referensi berhasil.|crud,reply
Referensi balasan tidak valid|Isi surat_masuk_id dengan ID 999999 yang tidak tersedia.|Penyimpanan ditolak.|Pesan validasi referensi surat masuk diterima.|reply
Perlindungan surat asal|Hapus atau reset surat masuk yang masih memiliki surat balasan.|Tindakan ditolak untuk menjaga hubungan surat.|Penghapusan dan reset ditolak; surat asal tetap tersedia.|protected
Reset tanpa konfirmasi|Reset surat keluar tanpa konfirmasi yang diwajibkan.|Reset ditolak.|Pesan validasi muncul dan data tetap ada.|reset
Reset dengan konfirmasi|Reset surat keluar menggunakan konfirmasi HAPUS.|Surat keluar dihapus; surat masuk tetap ada.|Daftar surat keluar kosong; data surat masuk tetap tersedia.|reset
Pratinjau impor CSV|Unggah CSV dengan 1 baris valid, 1 duplikat, dan 1 invalid.|Pratinjau membedakan ketiga status.|Jumlah pratinjau sesuai: 1 valid, 1 duplikat, 1 invalid.|csv
Konfirmasi impor valid|Simpan hasil pratinjau CSV tersebut.|Hanya baris valid ditambahkan.|Satu surat baru disimpan; duplikat dan baris invalid dilewati.|csv
Perlindungan konfirmasi impor|Gunakan ulang token impor atau konfirmasi dengan pengguna lain.|Konfirmasi yang tidak sah ditolak.|Token ulang dan token milik pengguna lain ditolak.|csv,badimport
Pemilihan Sheet1 dan Sheet2|Impor satu workbook: Sheet1 surat masuk dan Sheet2 surat keluar.|Masing-masing fitur membaca sheet yang sesuai.|Pratinjau dan konfirmasi kedua jenis surat berhasil memakai sheet yang sesuai.|sheets,b2
Impor agenda dengan header bertingkat|Gunakan agenda dengan header multiline dan tanggal pada sel gabungan.|Header dan tanggal agenda dikenali.|Data surat keluar valid dikenali dengan tanggal yang sesuai.|merged
Tanggal keluar kosong pada agenda lama|Kosongkan kolom tanggal keluar; isi tanggal surat pada agenda dengan tiga baris header.|Tanggal surat menjadi tanggal keluar sesuai aturan agenda lama; baris judul dilewati.|Dua data uji valid terbaca; tanggal terisi dan catatan penyesuaian tersedia.|fallback
Tanggal impor yang tidak valid|Isi tanggal agenda tidak valid, kosongkan kedua tanggal, atau kosongkan tanggal keluar pada template datar.|Input tanggal tidak valid tetap ditolak.|Ketiga kondisi menghasilkan baris invalid; tanggal salah tidak diganti otomatis.|invaliddate,bothdates,flatdate
Pemilihan sheet berdasarkan isi|Nama sheet tidak sesuai header; uji pula nama sheet yang cocok dengan huruf berbeda.|Isi header menentukan jenis; sheet bernama sesuai diprioritaskan bila isinya cocok.|Sheet yang tepat dipilih pada kedua kondisi.|sheetname,named
Workbook tanpa sheet surat keluar|Impor melalui fitur surat keluar dengan workbook yang hanya berisi surat masuk.|Pesan bahwa sheet sesuai tidak tersedia ditampilkan.|Impor ditolak; baris surat masuk tidak dibaca sebagai surat keluar.|nosheet
Format impor tidak sesuai|Unggah CSV dengan header not,a,valid,template.|Format tidak sesuai ditolak.|Pesan kesalahan format impor ditampilkan.|badimport
Ekspor surat ke Excel|Ekspor daftar surat masuk dan surat keluar.|Berkas Excel dapat diunduh.|Kedua endpoint memberi HTTP 200 dan berkas unduhan; unduhan browser berhasil.|export,b6
Perubahan profil akun|Ubah nama dan email dengan nilai valid.|Profil baru tersimpan.|Nama dan email diperbarui sesuai input.|profile
Validasi email profil dan admin|Gunakan email milik akun lain atau email berformat salah.|Email tidak valid atau duplikat ditolak.|Pesan validasi muncul; data profil lama tetap tersedia.|email
Perubahan kata sandi valid|Isi kata sandi lama benar, kata sandi baru, dan konfirmasi yang cocok.|Kata sandi baru tersimpan.|Kata sandi akun berhasil diperbarui dan cocok dengan nilai baru.|profile
Penolakan perubahan kata sandi|Gunakan kata sandi lama salah, kata sandi baru kurang dari 8 karakter, atau konfirmasi berbeda.|Perubahan ditolak dengan pesan terkait.|Ketiga kondisi menghasilkan kesalahan validasi.|password
Penambahan administrator|Isi nama, email unik, dan kata sandi valid.|Akun administrator baru tersedia.|Akun baru berhasil dibuat sesuai input.|profile
Penghapusan akun lain|Hapus administrator lain yang memiliki riwayat pencatatan surat.|Akun dinonaktifkan dan sesinya dicabut; riwayat pencatat terjaga.|Akun tidak dapat login; sesi dicabut dan identitas pencatat tetap ada.|account,deleted
Perlindungan akun sendiri|Administrator mencoba menghapus akunnya sendiri.|Penghapusan ditolak.|Akun aktif sendiri tetap tersedia.|account
Perlindungan identitas pencatat|Kirim created_by atau updated_by palsu saat menyimpan surat.|Input identitas palsu ditolak dan surat tidak disimpan.|Kesalahan validasi identitas muncul; tidak ada surat baru tersimpan.|forge
Tampilan dan menu ponsel|Buka detail balasan pada layar lebar 390 px lalu buka dan tutup menu.|Menu dapat digunakan tanpa halaman meluber horizontal.|Menu terbuka dan tertutup; tidak ada overflow horizontal.|b5
""".strip()

rows = []
for no, line in enumerate(data.splitlines(), 1):
    feature, action, expected, actual, refs = line.split('|')
    verified = []
    for key in refs.split(','):
        if key.startswith('b') and key[1:].isdigit():
            verified.append(browser['passed'][int(key[1:])])
        else:
            name = 'test_' + keys[key]
            assert name in cases, name
            verified.append(name)
    rows.append(dict(no=no, feature=feature, action=action, expected=expected, actual=actual, status='Lulus', evidence=verified))
assert len(rows) == 54

OUT = ROOT / 'docs'
OUT.mkdir(exist_ok=True)
(OUT / 'bukti-pengujian-blackbox.json').write_text(json.dumps(dict(date='2026-09-17', automated_tests=46, assertions=401, browser_groups=8, scenarios=rows), ensure_ascii=False, indent=2), encoding='utf-8')

intro = ('Pengujian dilaksanakan pada 17 September 2026 terhadap SIPAS berbasis Laravel dan MySQL 8.4.10. '
         'Pengujian menggunakan database khusus yang terpisah dari arsip operasional. Metode black box '
         'menilai kesesuaian input, keluaran, pesan validasi, dan perilaku fitur melalui permintaan HTTP dan browser Chrome.')
summary = ('Seluruh 54 skenario dalam tabel memenuhi hasil yang diharapkan. Bukti eksekusi keseluruhan: '
           '46 pengujian otomatis, 401 assertion, serta 8 kelompok pengujian browser lulus; tidak terdapat failure, '
           'error, atau permintaan aset lokal yang gagal pada eksekusi tersebut. Jumlah skenario tabel berbeda '
           'dari jumlah metode pengujian karena satu metode dapat memeriksa beberapa kondisi.')
notes = ('Status Lulus berarti hasil aktual sesuai hasil yang diharapkan, termasuk ketika sistem menolak input '
         'yang salah. Pengujian ini menilai fungsi pada lingkungan localhost dan tidak mencakup uji beban '
         'atau pengujian penetrasi. Hasil ini tidak menjamin seluruh kemungkinan kondisi bebas kesalahan.')
technical = ('Pemeriksaan teknis pendukung yang terpisah dari penilaian black box juga lulus: batasan foreign key MySQL, '
             'perlindungan penghapusan fisik pengguna yang direferensikan, rollback impor, penanganan kegagalan '
             'dan penggantian lampiran, serta pemindahan data SQLite. Build produksi Vite berhasil. '
             'Pemetaan setiap skenario terhadap bukti pengujian tersedia dalam bukti-pengujian-blackbox.json.')

md = ['# Tabel pengujian black box SIPAS', '', intro, '', summary, '',
      '| No | Fitur dan skenario | Input atau tindakan | Hasil yang diharapkan | Hasil aktual | Status |',
      '| --- | --- | --- | --- | --- | --- |']
for r in rows:
    md.append('| ' + ' | '.join(str(r[k]) for k in ['no','feature','action','expected','actual','status']) + ' |')
md += ['', notes, '', '## Pemeriksaan teknis pendukung', '', technical, '']
(OUT / 'tabel-pengujian-blackbox.md').write_text('\n'.join(md), encoding='utf-8')

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21.59), Cm(27.94)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name in ['Title','Heading 1','Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = __import__('docx').shared.RGBColor(0,0,0)
    for fonts in style.element.xpath('./w:rPr/w:rFonts'):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower(): del fonts.attrib[attr]
        fonts.set(qn('w:eastAsia'), 'Times New Roman')
        fonts.set(qn('w:cs'), 'Times New Roman')
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Title'].font.size = Pt(17)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(13)
doc.add_paragraph('Tabel Pengujian Black Box SIPAS', 'Title')
doc.add_paragraph(intro)
doc.add_paragraph(summary)
doc.add_paragraph('Tabel 1 Hasil pengujian fungsional sistem')

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [0.8, 4.4, 5.3, 5.6, 1.4]
for col, width in zip(table.columns, widths): col.width = Cm(width)
pr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
    elem = OxmlElement('w:'+edge)
    for key, value in [('val','single'),('sz','4'),('color','D9D9D9')]: elem.set(qn('w:'+key),value)
    borders.append(elem)
pr.append(borders)
headers = ['No','Fitur dan skenario','Hasil yang diharapkan','Hasil aktual','Status']
for cell, text in zip(table.rows[0].cells, headers): cell.text = text
repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
for r in rows:
    cells = table.add_row().cells
    cells[0].text = str(r['no'])
    p = cells[1].paragraphs[0]
    p.add_run(r['feature']).bold = True
    p.add_run('\n'+r['action'])
    cells[2].text, cells[3].text, cells[4].text = r['expected'], r['actual'], r['status']
for ri, row in enumerate(table.rows):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for ci, (cell, width) in enumerate(zip(row.cells, widths)):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            el = OxmlElement('w:'+edge); el.set(qn('w:w'),'80'); el.set(qn('w:type'),'dxa'); margins.append(el)
        tcpr.append(margins)
        if ri == 0:
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'F2F2F2'); tcpr.append(shade)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in [0,4] else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name, run.font.size = 'Times New Roman', Pt(11)
                if ri == 0: run.bold = True
doc.add_paragraph(notes).paragraph_format.space_before = Pt(6)
doc.add_paragraph('Pemeriksaan teknis pendukung', 'Heading 1').paragraph_format.space_before = Pt(12)
doc.add_paragraph(technical)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Halaman ')
for instruction in ['PAGE','NUMPAGES']:
    if instruction == 'NUMPAGES': footer.add_run(' dari ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), instruction); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(10)
doc.save(OUT / 'Tabel_Pengujian_BlackBox_SIPAS.docx')
print('Created DOCX, Markdown, and evidence mapping for 54 verified scenarios.')
