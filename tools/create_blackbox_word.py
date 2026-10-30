from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
root=Path.cwd(); data=json.loads((root/'docs/bukti-pengujian-blackbox.json').read_text(encoding='utf-8'))
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.6)
for n in ['Normal','Title','Heading 1']:
 st=doc.styles[n]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
for st in doc.styles:
 for borders in list(st.element.iter(qn('w:pBdr'))): borders.getparent().remove(borders)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(16)
doc.add_paragraph('Pengujian Black Box Aplikasi Sistem Surat',style='Title')
doc.add_paragraph('Pengujian black box dilakukan untuk memeriksa kesesuaian fungsi aplikasi berdasarkan masukan dan keluaran yang dihasilkan. Tabel berikut memuat 54 skenario yang mencakup autentikasi, pengelolaan surat, lampiran, pencarian, impor, ekspor, serta pengelolaan akun.')
doc.add_paragraph('Hasil aktual mengacu pada bukti pengujian tanggal 17 September 2026 melalui permintaan HTTP dan browser Chrome pada lingkungan localhost dengan database pengujian terpisah. Sebanyak 46 pengujian otomatis dengan 401 pemeriksaan dan 8 kelompok pengujian browser tercatat lulus. Satu pengujian dapat mencakup beberapa skenario tabel.')
p=doc.add_paragraph('Tabel Pengujian Black Box Aplikasi Sistem Surat'); p.runs[0].bold=True; p.paragraph_format.keep_with_next=True
headers=['No.','Skenario pengujian','Input atau tindakan','Hasil yang diharapkan','Hasil aktual','Status']
widths=[.35,1.25,1.5,1.55,2.05,.6]
t=doc.add_table(rows=1,cols=6); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for c,w in zip(t.columns,widths): c.width=Inches(w)
for c,w,h in zip(t.rows[0].cells,widths,headers):c.width=Inches(w);c.text=h
repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
for s in data['scenarios']:
 cells=t.add_row().cells
 values=[str(s['no']),s['feature'],s['action'],s['expected'],s['actual'],s['status']]
 for c,w,v in zip(cells,widths,values):c.width=Inches(w);c.text=v
for ri,r in enumerate(t.rows):
 trpr=r._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
 for ci,c in enumerate(r.cells):
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
  pr=c._tc.get_or_add_tcPr()
  mar=OxmlElement('w:tcMar')
  for name,value in [('top',90),('bottom',90),('left',70),('right',70)]:
   e=OxmlElement('w:'+name);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
  pr.append(mar)
  borders=OxmlElement('w:tcBorders')
  for name in ['top','left','bottom','right']:
   e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'808080');borders.append(e)
  pr.append(borders)
  if ri==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');pr.append(sh)
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
   if ci in [0,5]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for run in p.runs:run.font.size=Pt(10);run.bold=ri==0;run.font.name='Times New Roman'
doc.add_paragraph('Status Lulus menunjukkan bahwa hasil aktual sesuai dengan hasil yang diharapkan, termasuk penolakan terhadap masukan yang tidak valid. Seluruh 54 skenario tercatat lulus pada pengujian tersebut. Cakupan pengujian berfokus pada fungsi aplikasi dan tidak mencakup uji beban atau pengujian penetrasi.')
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Halaman ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
out=root/'output';out.mkdir(exist_ok=True)
file=out/'Tabel-Pengujian-Blackbox-Sistem-Surat.docx';doc.save(file)
print(file)

