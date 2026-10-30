"""Revise the supplied PKL report while preserving its school template."""
from copy import deepcopy
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r'C:/Users/zahra/Pictures/LAPORAN PKL/Laporan PKL_Dinas Paternakan.docx')
DEST = SOURCE.with_name('Laporan PKL_Dinas Paternakan_Revisi MySQL.docx')
ASSETS = ROOT / '.local/laporan-mysql-assets'
ASSETS.mkdir(exist_ok=True)
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
schema = json.loads((ROOT / 'database/schema-mysql.json').read_text(encoding='utf-8'))
tables, relations = schema['tables'], schema['relations']
assert len(tables) == 11 and len(relations) == 6 and schema['version'] == '8.4.10'

def font(size, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'), size)

def structure(name):
    cols = tables[name]
    rels = [r for r in relations if r['child_table'] == name]
    height = 122 + len(cols)*43 + 62 + len(rels)*36
    im = Image.new('RGB', (1320,height),'white'); draw = ImageDraw.Draw(im)
    draw.text((24,16),f'Struktur tabel {name}',font=font(35,True),fill='black')
    draw.text((24,65),'MySQL 8.4.10   Database sipas_db   Sumber skema aktual',font=font(25),fill='#444444')
    positions = [24,560,1030,1180]
    headings = ['Kolom','Tipe MySQL','NULL','Key']
    draw.rectangle((15,102,1305,143),fill='#eeeeee',outline='#aaaaaa')
    for x,h in zip(positions,headings): draw.text((x,108),h,font=font(28,True),fill='black')
    for i,c in enumerate(cols):
        y=143+i*43
        draw.rectangle((15,y,1305,y+43),fill='#f8f8f8' if i%2 else 'white',outline='#cccccc')
        fk=any(r['child_column']==c['name'] for r in rels)
        key='PK' if c['column_key']=='PRI' else ('FK' if fk else ('UK' if c['column_key']=='UNI' else '-'))
        values=[c['name'],c['type'],'Ya' if c['nullable']=='YES' else 'Tidak',key]
        for x,value in zip(positions,values): draw.text((x,y+7),value,font=font(28,key=='PK'),fill='black')
    y=151+len(cols)*43
    draw.text((24,y),'PK kunci utama   UK nilai unik   FK kunci asing',font=font(25),fill='#333333')
    for i,r in enumerate(rels):
        draw.text((24,y+37+i*36),f"FK {r['child_column']} -> {r['parent_table']}.{r['parent_column']}   DELETE {r['delete_rule']}",font=font(25),fill='black')
    path=ASSETS/f'{name}.png'; im.save(path); return path

def overview():
    im=Image.new('RGB',(1320,720),'white'); draw=ImageDraw.Draw(im)
    draw.text((28,22),'Basis data SIPAS MySQL',font=font(42,True),fill='black')
    draw.text((28,85),'MySQL 8.4.10   sipas_db   11 tabel   6 foreign key',font=font(32),fill='black')
    draw.rectangle((24,147,1295,208),fill='#eeeeee',outline='#aaaaaa')
    draw.text((44,158),'3 tabel utama',font=font(33,True),fill='black')
    draw.text((690,158),'8 tabel pendukung Laravel',font=font(33,True),fill='black')
    for i,name in enumerate(['users','surat_masuk','surat_keluar']):
        draw.text((44,235+i*70),name,font=font(35),fill='black')
    support=['sessions','migrations','password_reset_tokens','cache','cache_locks','jobs','job_batches','failed_jobs']
    for i,name in enumerate(support): draw.text((690,232+i*52),name,font=font(32),fill='black')
    draw.text((44,490),'5 FK pada tabel surat',font=font(30),fill='black')
    draw.text((44,540),'1 FK pada tabel sessions',font=font(30),fill='black')
    draw.text((44,638),'Sumber information_schema dan SHOW CREATE TABLE pada sipas_db',font=font(27),fill='#444444')
    path=ASSETS/'database-mysql.png'; im.save(path); return path

def erd():
    im=Image.new('RGB',(1900,1740),'white'); draw=ImageDraw.Draw(im)
    draw.text((65,24),'ERD SIPAS MySQL dengan relasi surat dan sesi',font=font(44,True),fill='black')
    locations={'users':(690,130),'sessions':(70,130),'surat_masuk':(70,710),'surat_keluar':(1290,710)}
    row_h,header,w=45,60,540
    def y(name,col): return locations[name][1]+header+row_h*next(i for i,c in enumerate(tables[name]) if c['name']==col)+row_h/2
    def wire(points): draw.line(points,fill='#555555',width=4)
    uy=y('users','id')
    for field,channel in [('created_by',0),('updated_by',1)]:
        ty=y('surat_masuk',field); x=645-30*channel
        wire([(690,uy),(x,uy),(x,ty),(610,ty)])
        ty=y('surat_keluar',field); x=1250+25*channel
        wire([(1230,uy),(x,uy),(x,ty),(1290,ty)])
    sy=y('sessions','user_id')
    wire([(690,uy),(630,uy),(630,sy),(610,sy)])
    iy,oy=y('surat_masuk','id'),y('surat_keluar','surat_masuk_id')
    wire([(70,iy),(25,iy),(25,1580),(1870,1580),(1870,oy),(1830,oy)])
    for name,(x,top) in locations.items():
        draw.rectangle((x,top,x+w,top+header+row_h*len(tables[name])),fill='white',outline='#666666',width=3)
        draw.rectangle((x,top,x+w,top+header),fill='#eeeeee',outline='#666666',width=3)
        draw.text((x+20,top+10),name,font=font(39,True),fill='black')
        for i,c in enumerate(tables[name]):
            yy=top+header+i*row_h
            draw.line((x,yy,x+w,yy),fill='#cccccc',width=2)
            fk=any(r['child_table']==name and r['child_column']==c['name'] for r in relations)
            key='PK' if c['column_key']=='PRI' else ('FK' if fk else '')
            draw.text((x+10,yy+7),key,font=font(28,True),fill='black')
            draw.text((x+68,yy+5),c['name'],font=font(32,key=='PK'),fill='black')
    draw.text((612,sy-34),'0..N',font=font(24,True),fill='black')
    for field in ['created_by','updated_by']:
        draw.text((612,y('surat_masuk',field)-33),'0..N',font=font(24,True),fill='black')
        draw.text((1224,y('surat_keluar',field)-33),'0..N',font=font(24,True),fill='black')
    draw.text((7,iy-42),'0..1',font=font(24,True),fill='black')
    draw.text((1835,oy-38),'0..N',font=font(24,True),fill='black')
    draw.text((708,645),'Setiap FK menunjuk 0..1 induk',font=font(30),fill='black')
    draw.text((708,690),'karena kolom FK nullable',font=font(30),fill='black')
    notes=['created_by administrator pencatat','updated_by pengubah terakhir','surat_masuk_id asal surat balasan','Lima FK surat memakai RESTRICT','FK sesi memakai CASCADE']
    for i,note in enumerate(notes): draw.text((695,880+i*65),note,font=font(29),fill='black')
    draw.text((65,1620),'PK kunci utama   FK kunci asing   0..N nol sampai banyak data anak',font=font(34),fill='black')
    draw.text((65,1672),'Sumber skema aktual sipas_db   Tabel pendukung selain sessions tidak ditampilkan',font=font(31),fill='#444444')
    path=ASSETS/'erd-mysql-laporan.png'; im.save(path); return path

figures={11:erd(),12:overview(),13:structure('users'),14:structure('sessions'),15:structure('surat_masuk'),16:structure('surat_keluar')}
doc=Document(SOURCE)
ps=list(doc.paragraphs)
changes=[]

replacements={
84:'Beberapa penelitian sebelumnya menunjukkan bahwa sistem informasi berbasis web dapat membantu pengelolaan administrasi surat. Mayu dan Mira mengembangkan sistem pencatatan surat berbasis website untuk mendukung pencatatan serta penelusuran surat [3]. Azdiansyah dan Azhar mengembangkan sistem pengelolaan surat pada instansi pemerintah yang mencakup pencatatan digital, pencarian arsip, dan pelaporan; hasil evaluasinya menunjukkan bahwa fitur inti dapat digunakan sesuai kebutuhan pengguna [4].',
140:'Peternakan merupakan usaha pengelolaan dan pemeliharaan ternak untuk menghasilkan produk peternakan. Kegiatan tersebut menghasilkan data yang perlu dikelola, antara lain data ternak, pakan, kematian, pembelian, penjualan, serta biaya pemeliharaan. Murniawati et al. (2024) menjelaskan kebutuhan pengelolaan data tersebut melalui penelitian sistem informasi peternakan ayam [8].',
141:'Pengelolaan data yang terstruktur membantu pencarian dan penyajian informasi peternakan. Murniawati et al. menemukan bahwa pencatatan manual menyebabkan data sulit ditelusuri, kemudian mengembangkan sistem berbasis web untuk membantu pengelolaan data dan penyusunan laporan [8].',
147:'Pengelolaan surat masuk dan surat keluar secara manual sering menimbulkan kendala karena data tersimpan dalam catatan fisik atau arsip kertas. Praja et al. menjelaskan bahwa pencatatan pada buku agenda dan penyimpanan fisik menyulitkan pencarian arsip [2]. Mayu dan Mira mengembangkan sistem pencatatan berbasis website untuk mendukung pengelolaan surat pada instansi [3].',
150:'Arsip digital merupakan pengelolaan dokumen dengan teknologi informasi sehingga data dan berkas dapat disimpan serta ditelusuri melalui sistem komputer. Azdiansyah dan Azhar menunjukkan penerapan pencatatan digital, pencarian, dan pengarsipan pada sistem persuratan instansi pemerintah [4]. Candra et al. menjelaskan bahwa pengarsipan digital membantu pengelolaan dan pencarian data surat [7].',
153:'Pada administrasi surat, aplikasi berbasis web mendukung pencatatan, pencarian, perubahan, penghapusan, serta penyajian laporan. Purnomo et al. mengembangkan sistem administrasi surat berbasis web untuk mendukung pengelolaan dan penyimpanan arsip digital [5]. Londa et al. menerapkan sistem berbasis website untuk memantau disposisi surat masuk dan surat keluar [6].',
156:'Laravel merupakan framework open-source berbasis PHP untuk membangun aplikasi web dengan struktur yang terorganisasi. Widyastuti et al. (2024) menggunakan Laravel dalam pengembangan sistem informasi perpustakaan berbasis web [9].',
157:'Laravel mendukung pengembangan pengelolaan data, autentikasi, dan modul aplikasi. Fasilitas tersebut digunakan untuk menghubungkan antarmuka pengguna dengan proses pengolahan data pada sistem [9].',
158:'Tahapan penelitian Widyastuti et al. meliputi analisis kebutuhan, perancangan, pengembangan, pengujian, serta evaluasi. Sistem yang dihasilkan membantu pengelolaan informasi perpustakaan yang sebelumnya dilakukan secara manual [9].',
159:'Hasil penelitian tersebut menunjukkan perbaikan kemudahan penggunaan, kecepatan akses, keamanan data, dan efisiensi operasional pada sistem yang diuji [9].',
161:'Database MySQL',
162:'Basis data merupakan kumpulan data yang disusun secara terstruktur untuk mendukung penyimpanan, pengolahan, dan penelusuran informasi. Siregar et al. (2024) menjelaskan bahwa Database Management System (DBMS) menyediakan mekanisme pengelolaan data, sedangkan Relational Database Management System (RDBMS) mendukung hubungan antartabel [10].',
163:'MySQL merupakan RDBMS yang menggunakan Structured Query Language (SQL) untuk mengakses dan mengolah data. MySQL dijalankan sebagai server basis data yang menerima permintaan dari aplikasi klien. Ratnasari et al. (2024) menerapkan MySQL sebagai penyimpanan terpusat pada sistem client-server, dengan aplikasi web memproses masukan pengguna sebelum menyimpan dan menampilkan data [11].',
164:'Perancangan basis data MySQL mencakup penentuan tabel, kolom, tipe data, primary key, dan foreign key. Primary key mengidentifikasi setiap baris, sedangkan foreign key menghubungkan data anak dengan data induk. Siregar et al. menekankan penetapan hubungan antartabel untuk menjaga integritas data serta penggunaan SQL melalui operasi SELECT, INSERT, UPDATE, dan DELETE [10].',
165:'Penggunaan MySQL pada administrasi surat juga diterapkan oleh Iman dan Supardi (2025) dalam sistem arsip surat masuk dan surat keluar pada Kantor Camat Labuhan Deli. Penelitian tersebut menggunakan PHP dan MySQL serta menunjukkan bahwa sistem terkomputerisasi dapat membantu pengarsipan dan pemantauan surat [12]. Penelitian ini menjadi contoh penerapan MySQL pada kebutuhan administrasi yang sejenis.',
166:'Pada versi akhir Sistem Informasi Pengelolaan Administrasi Surat (SIPAS), MySQL 8.4.10 digunakan melalui database sipas_db. Browser berkomunikasi dengan aplikasi Laravel melalui HTTP; Laravel memvalidasi masukan dan menjalankan operasi basis data melalui koneksi MySQL. Tiga tabel utama adalah users, surat_masuk, dan surat_keluar. Delapan tabel pendukung digunakan oleh Laravel, termasuk sessions. Enam foreign key diterapkan untuk menghubungkan pencatat dan pengubah surat, referensi surat balasan, serta sesi pengguna. Informasi lampiran disimpan pada tabel surat, sedangkan berkas fisiknya disimpan pada storage aplikasi.',
168:'Import data merupakan proses memasukkan data dari file Excel ke basis data melalui tahapan unggah, validasi, pembersihan, dan pratinjau. Putri et al. (2026) mengembangkan sistem konversi Excel ke MySQL yang memeriksa duplikat, kelengkapan, serta kesesuaian format sebelum data disimpan [13]. Pada SIPAS, pratinjau membedakan baris valid, duplikat, dan invalid; hanya baris valid yang disimpan setelah konfirmasi pengguna.',
169:'Export data merupakan proses mengeluarkan data sistem ke format Excel untuk kebutuhan agenda dan pelaporan. Pada SIPAS, data dibaca dari MySQL berdasarkan pencarian serta filter bulan dan tahun, kemudian disusun menjadi file Excel. Dengan demikian, fitur import dan export menghubungkan arsip spreadsheet dengan penyimpanan terpusat pada aplikasi.',
176:'Ruang lingkup pekerjaan meliputi analisis kebutuhan sistem, perancangan alur, perancangan basis data, implementasi aplikasi, serta pengujian fitur. Sistem dibangun menggunakan Laravel dan PHP. Hasil penyempurnaan pada versi akhir sistem menggunakan MySQL 8.4.10 sebagai basis data, dengan relasi antartabel diterapkan melalui foreign key untuk menjaga konsistensi identitas pencatat dan keterkaitan surat balasan.',
197:'Use Case Diagram menggambarkan interaksi admin dengan fitur dashboard, surat masuk, surat keluar, profil, dan logout. Perancangan basis data menggunakan tiga tabel utama, yaitu users untuk akun administrator, surat_masuk untuk arsip surat diterima, serta surat_keluar untuk arsip surat dikirim. Tabel sessions merupakan tabel pendukung autentikasi. Kolom created_by dan updated_by pada kedua tabel surat merujuk users.id, surat_keluar.surat_masuk_id merujuk surat_masuk.id, dan sessions.user_id merujuk users.id. Relasi tersebut diterapkan sebagai foreign key pada MySQL.',
205:'Fitur surat keluar mengelola tanggal keluar, nomor agenda, alamat penerima, tanggal surat, nomor surat, perihal, asal surat, dan lampiran. Jika surat keluar merupakan balasan, admin dapat memilih surat masuk asal pada formulir. Pilihan tersebut bersifat opsional. Identitas pencatat dan pengubah terakhir ditentukan berdasarkan akun yang sedang login.',
268:'Pada fitur surat masuk, admin dapat melihat daftar, menambah, mengubah, menghapus, mengunggah lampiran, mencari, serta memfilter surat berdasarkan periode. Impor Excel/CSV dilakukan melalui pratinjau dan konfirmasi sebelum baris valid disimpan ke MySQL. Admin juga dapat mengekspor data dan melakukan reset dengan konfirmasi. Penghapusan atau reset surat masuk ditolak apabila data yang akan dihapus masih memiliki surat keluar sebagai balasan.',
269:'Pada fitur surat keluar, admin dapat melihat daftar, menambah, mengubah, menghapus, mengelola lampiran, mencari, dan memfilter data. Admin dapat mengimpor melalui pratinjau, mengekspor Excel, serta melakukan reset dengan konfirmasi. Surat keluar dapat dihubungkan ke surat masuk asal sebagai balasan. Detail kedua surat menyediakan tautan untuk menelusuri hubungan tersebut.',
270:'Pada pengelolaan profil, admin dapat mengubah nama dan email, mengganti kata sandi, menambah administrator, serta menonaktifkan akun administrator lain. Penonaktifan menggunakan soft delete agar identitas pencatat pada riwayat surat tetap tersedia. Sesi akun yang dinonaktifkan dicabut. Sistem melindungi akun sendiri dan administrator aktif terakhir dari penghapusan.',
273:'Entity Relationship Diagram (ERD) menggambarkan tabel beserta hubungan data pada MySQL. Diagram berikut menampilkan tiga tabel utama dan tabel sessions sebagai pendukung autentikasi. Hubungan yang digambarkan sesuai dengan enam foreign key pada database sipas_db.',
275:'Pada Gambar 5.2, ERD menampilkan users, surat_masuk, surat_keluar, dan sessions. Tabel users menyimpan akun administrator serta deleted_at untuk penonaktifan akun. Setiap pengguna dapat menjadi pencatat maupun pengubah banyak surat melalui relasi terpisah pada created_by dan updated_by. Setiap sesi terautentikasi dihubungkan ke users melalui sessions.user_id, sedangkan sesi tamu dapat memiliki user_id kosong.',
276:'Tabel surat_masuk menyimpan tanggal penerimaan, nomor agenda, pengirim, tanggal dan nomor surat, perihal, tujuan disposisi, serta metadata lampiran. Kolom created_by merujuk administrator pencatat dan updated_by merujuk administrator yang terakhir mengubah data. Kedua kolom menggunakan foreign key ke users.id dengan aturan ON DELETE RESTRICT.',
277:'Tabel surat_keluar menyimpan tanggal pengiriman, nomor agenda, penerima, tanggal dan nomor surat, perihal, asal surat, serta metadata lampiran. Kolom created_by dan updated_by merujuk users.id. Kolom surat_masuk_id merujuk surat_masuk.id dengan aturan ON DELETE RESTRICT. Satu surat masuk dapat memiliki nol atau banyak surat balasan, sedangkan satu surat keluar merujuk nol atau satu surat masuk asal.',
278:'Seluruh kolom foreign key bersifat nullable untuk mengakomodasi data lama, sesi tamu, atau hubungan balasan yang tidak digunakan. Aplikasi tetap mengisi pencatat untuk surat baru berdasarkan pengguna yang login. Lima foreign key pada tabel surat memakai RESTRICT untuk mencegah penghapusan fisik data induk yang masih dirujuk; foreign key sesi memakai ON DELETE CASCADE. Penghapusan akun melalui aplikasi menggunakan soft delete serta pencabutan sesi, sehingga identitas pencatat pada arsip tetap dapat ditelusuri.',
280:'Tahap Coding menerapkan rancangan ke dalam aplikasi Laravel dengan bahasa PHP. Pada versi akhir sistem, penyimpanan data menggunakan MySQL 8.4.10. Struktur tabel, indeks, dan foreign key dibentuk melalui migration; relasi yang sama didefinisikan pada model Eloquent agar data pencatat, pengubah, sesi, dan surat balasan dapat diakses melalui aplikasi.',
283:'Implementasi basis data menggunakan MySQL 8.4.10 dengan database sipas_db. Skema terdiri atas 11 tabel, yaitu tiga tabel utama dan delapan tabel pendukung Laravel. Tabel utama adalah users, surat_masuk, dan surat_keluar. Tabel pendukung meliputi sessions, migrations, password_reset_tokens, cache, cache_locks, jobs, job_batches, dan failed_jobs. Enam foreign key diterapkan sesuai ERD. Penyimpanan surat dan konfirmasi impor menggunakan transaksi, sehingga kegagalan proses dapat membatalkan perubahan data. Visualisasi struktur berikut disusun dari skema MySQL aktual.',
287:'Berdasarkan Gambar 5.4, tabel users menyimpan identitas dan kredensial administrator. Kolom id bertipe BIGINT UNSIGNED sebagai primary key, sedangkan email memiliki batasan unik. Kolom password menyimpan hash kata sandi dan remember_token mendukung fitur mengingat login. Kolom email_verified_at, created_at, dan updated_at menyimpan informasi waktu. Kolom deleted_at mendukung soft delete; akun yang dinonaktifkan tidak dapat login, tetapi tetap tersedia sebagai referensi pencatat pada riwayat surat.',
289:'Berdasarkan Gambar 5.5, tabel sessions menyimpan id sesi, user_id, ip_address, user_agent, payload, serta last_activity. Kolom user_id bertipe BIGINT UNSIGNED, dapat kosong untuk sesi tamu, dan memiliki foreign key ke users.id dengan ON DELETE CASCADE. Satu pengguna dapat memiliki banyak sesi. Saat akun dinonaktifkan melalui aplikasi, sesi akun tersebut dicabut secara eksplisit agar akses tidak dapat dilanjutkan.',
291:'Berdasarkan Gambar 5.6, tabel surat_masuk menyimpan tanggal penerimaan, nomor agenda, alamat pengirim, tanggal dan nomor surat, perihal, serta tujuan disposisi. Primary key id bertipe BIGINT UNSIGNED. Metadata lampiran dicatat pada lampiran_path, lampiran_nama_asli, lampiran_mime, dan lampiran_ukuran; berkas fisik tersimpan pada storage aplikasi. Kolom created_at dan updated_at mencatat waktu perubahan. Kolom created_by dan updated_by bertipe BIGINT UNSIGNED serta memiliki foreign key ke users.id dengan ON DELETE RESTRICT.',
293:'Berdasarkan Gambar 5.7, tabel surat_keluar menyimpan tanggal pengiriman, nomor agenda, alamat penerima, tanggal dan nomor surat, perihal, asal surat, serta metadata lampiran. Kolom created_by dan updated_by merujuk users.id. Kolom surat_masuk_id merujuk surat_masuk.id dan dapat kosong jika surat keluar bukan balasan. Ketiga foreign key menggunakan ON DELETE RESTRICT. Struktur tersebut mendukung penelusuran pencatat, pengubah terakhir, dan surat masuk asal balasan.',
295:'Setelah struktur MySQL diterapkan, fungsi sistem dihubungkan dengan antarmuka aplikasi. Laravel menerima permintaan pengguna, memeriksa autentikasi dan validasi, lalu membaca atau menyimpan data melalui koneksi MySQL. Hasilnya ditampilkan pada halaman berikut.',
297:'Berdasarkan Gambar 5.8, halaman login digunakan administrator untuk mengakses sistem melalui email dan kata sandi yang terdaftar. Laravel memeriksa kesesuaian kredensial serta status aktif akun pada tabel users. Jika autentikasi berhasil, pengguna diarahkan ke dashboard dan sesi disimpan pada tabel sessions. Jika kredensial salah atau akun telah dinonaktifkan, sistem menolak login dengan pesan kesalahan.',
300:'Berdasarkan Gambar 5.9, dashboard menampilkan jumlah surat masuk dan surat keluar pada bulan berjalan, jumlah surat tahun berjalan, total seluruh surat, grafik bulanan, serta aktivitas terbaru. Data dihitung dari tabel surat pada MySQL. Dashboard membantu administrator memantau ringkasan arsip dan mengakses menu pengelolaan surat.',
302:'Berdasarkan Gambar 5.10, halaman surat masuk mengelola tanggal penerimaan, nomor agenda, pengirim, tanggal dan nomor surat, perihal, tujuan disposisi, serta lampiran. Administrator dapat menambah, melihat, mengubah, menghapus, mencari, dan memfilter surat. Impor Excel/CSV memakai pratinjau dan konfirmasi, sedangkan ekspor menghasilkan Excel berdasarkan filter. Detail surat menampilkan pencatat, pengubah terakhir, dan daftar surat balasan. Surat masuk yang masih dirujuk balasan dilindungi dari penghapusan dan reset.',
304:'Berdasarkan Gambar 5.11, halaman surat keluar mengelola tanggal pengiriman, nomor agenda, penerima, tanggal dan nomor surat, perihal, asal, serta lampiran. Halaman menyediakan pengelolaan data, pencarian, filter, impor dengan pratinjau, dan ekspor Excel. Surat keluar yang merupakan balasan dapat dihubungkan ke surat masuk asal melalui formulir. Referensi tersebut disimpan pada surat_masuk_id dan ditampilkan sebagai tautan pada halaman detail.',
306:'Berdasarkan Gambar 5.12, halaman profil digunakan untuk memperbarui nama dan email, mengganti kata sandi dengan verifikasi kata sandi lama, serta menambah administrator. Penghapusan akun administrator lain dilakukan melalui soft delete pada users.deleted_at dan pencabutan sesi. Sistem melindungi akun sendiri serta administrator aktif terakhir. Riwayat surat tetap menampilkan identitas pencatat meskipun akun telah dinonaktifkan.',
}

def set_text(p,text,bold=False):
    p.clear()
    run=p.add_run(text)
    run.font.name='Times New Roman'; run.font.size=Pt(12); run.bold=bold; run.font.color.rgb=RGBColor(0,0,0)
    if p._p.pPr is not None:
        for elem in p._p.pPr.xpath('./w:shd|./w:rPr/w:highlight|./w:rPr/w:shd'): elem.getparent().remove(elem)

for i,new in replacements.items():
    changes.append({'paragraph':i,'before':ps[i].text,'after':new})
    set_text(ps[i],new,bold=i==161)

# Keep earlier numerical citations aligned with the repaired bibliography.
set_text(ps[143],ps[143].text)

captions={284:'Gambar 5.3 Basis Data MySQL',286:'Gambar 5.4 Struktur Tabel Users',288:'Gambar 5.5 Struktur Tabel Sessions',290:'Gambar 5.6 Struktur Tabel Surat Masuk',292:'Gambar 5.7 Struktur Tabel Surat Keluar',296:'Gambar 5.8 Halaman Login',299:'Gambar 5.9 Halaman Dashboard',301:'Gambar 5.10 Halaman Data Surat Masuk',303:'Gambar 5.11 Halaman Data Surat Keluar',305:'Gambar 5.12 Halaman Profil Akun'}
for i,text in captions.items():
    changes.append({'paragraph':i,'before':ps[i].text,'after':text})
    set_text(ps[i],text)
    ps[i].alignment=WD_ALIGN_PARAGRAPH.CENTER
    ps[i].paragraph_format.keep_with_next=False

# Replace outdated SQLite screenshots with visualizations of the actual MySQL schema.
for ti,image in figures.items():
    table=doc.tables[ti]
    cell=table.cell(0,0)
    for child in list(cell._tc):
        if child.tag!=qn('w:tcPr'): cell._tc.remove(child)
    pp=cell.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_before=Pt(3); pp.paragraph_format.space_after=Pt(3)
    pp.paragraph_format.line_spacing=1
    pp.paragraph_format.keep_with_next=True
    width = 14.5 if ti==11 else (13.3 if ti==15 else 14.8)
    pp.add_run().add_picture(str(image),width=Cm(width))
    rowpr=table.rows[0]._tr.get_or_add_trPr()
    for h in rowpr.xpath('./w:trHeight'): rowpr.remove(h)
    rowpr.append(OxmlElement('w:cantSplit'))

refs=[
'M. Z. I. Mansis, M. Al Fayed, and Irwan, “Perancangan Sistem Informasi Administrasi Surat Masuk dan Surat Keluar Pada Dinas Kominfo Serdang Bedagai,” Jurnal Minfo Polgan, vol. 13, no. 2, pp. 2049–2055, 2024, doi: 10.33395/jmp.v13i2.14375.',
'A. B. Praja, Darmansah, and S. Wijayanto, “Sistem Informasi Pencatatan Surat Masuk dan Surat Keluar Berbasis Website Menggunakan Metode Waterfall,” Jurnal Sistem Komputer dan Informatika (JSON), vol. 3, no. 3, pp. 273–282, 2022, doi: 10.30865/json.v3i3.3914.',
'I. Mayu and Mira, “Perancangan Sistem Informasi Pencatatan Surat Masuk dan Surat Keluar di Desa Tirta Kencana Berbasis Website,” Jurnal InFact Sains dan Komputer, vol. 7, no. 2, pp. 58–61, 2023, doi: 10.61179/jurnalinfact.v7i02.443.',
'M. Azdiansyah and N. C. Azhar, “Pengembangan Sistem Informasi Surat Masuk dan Keluar Berbasis Web pada Instansi Pemerintah dengan Evaluasi UAT,” EXPERT: Jurnal Manajemen Sistem Informasi dan Teknologi, vol. 15, no. 1, pp. 82–90, 2025, doi: 10.36448/expert.v15i1.4300.',
'D. R. Purnomo, R. Hamid, and I. D. Kurniawati, “Rancang Bangun Sistem Administrasi Surat Keluar Masuk Berbasis Web di Kantor Kecamatan Margomulyo,” Bulletin Ilmiah Informatika Teknologi, vol. 1, no. 2, pp. 52–58, 2022.',
'M. A. Londa, Y. A. Wee, and M. Radja, “Implementasi Sistem Informasi Monitoring Disposisi Surat Masuk dan Surat Keluar Berbasis Website,” MATRIK: Jurnal Manajemen, Teknik Informatika dan Rekayasa Komputer, vol. 21, no. 2, pp. 379–388, 2022, doi: 10.30812/matrik.v21i2.1443.',
'B. P. Candra, I. Thoib, P. R. Maulidiah, and R. Khamidah, “Perancangan Sistem Informasi Arsip Surat Masuk dan Keluar dengan Metode Waterfall,” Digital Transformation Technology, vol. 4, no. 2, pp. 959–966, 2024, doi: 10.47709/digitech.v4i2.5120.',
'M. Murniawati, A. Susanto, and A. A. Riadi, “Sistem Informasi Pengelolaan Peternakan Ayam,” JATI (Jurnal Mahasiswa Teknik Informatika), vol. 8, no. 1, pp. 1200–1206, 2024, doi: 10.36040/jati.v8i1.8953.',
'I. Widyastuti, M. H. Harike, M. N. Takbir, A. Malik, and J. Y. Sari, “Digital Transformation of Libraries: Web-based Information System Development with Laravel,” Journal of Embedded Systems, Security and Intelligent Systems, vol. 5, no. 2, pp. 147–152, 2024, doi: 10.59562/jessi.v5i2.3330.',
'U. K. Siregar, T. A. Sitakar, S. Haramain, Z. N. S. Lubis, U. Nadhirah, and Yahfizham, “Pengembangan database Management system menggunakan My SQL,” Jurnal Sains, Teknologi & Komputer, vol. 1, no. 1, pp. 8–12, 2024, doi: 10.56495/saintek.v1i1.450.',
'Ratnasari, N. Hafiizah, N. W. Erdiani, and A. H. Wibowo, “Implementasi Sistem Client-Server Berbasis Flask dan MySQL untuk Pengelolaan Data Mahasiswa dengan Protokol TCP/IP,” JATI (Jurnal Mahasiswa Teknik Informatika), vol. 8, no. 6, pp. 12577–12582, 2024, doi: 10.36040/jati.v8i6.11999.',
'H. R. Iman and Supardi, “Perancangan Sistem Informasi Arsip Surat Masuk dan Surat Keluar pada Kantor Camat Labuhan Deli,” Jurnal Sistem Informasi Komputer (SIKOM), vol. 2, no. 3, pp. 92–100, 2025. [Online]. Available: https://journal.beaninstitute.id/index.php/sikom/article/view/114.',
'C. N. Putri, F. Ilham, M. Hatun S, and T. V. Aulia, “Rancang Bangun Sistem Informasi Data Cleansing dan Konversi Data Excel ke Database MySQL Berbasis Web di SMK Makarya 1 Jakarta,” LOGIC: Jurnal Ilmu Komputer dan Pendidikan, vol. 4, no. 3, pp. 111–121, 2026. [Online]. Available: https://www.journal.mediapublikasi.id/index.php/logic/article/view/6248.',
]
for i,text in enumerate(refs[:9],374):
    set_text(ps[i],text)
    ps[i].paragraph_format.keep_together=True
    ps[i].paragraph_format.keep_with_next=False
ps[382].paragraph_format.page_break_before=True
anchor=ps[382]
for text in refs[9:]:
    el=OxmlElement('w:p'); el.append(deepcopy(ps[382]._p.pPr)); anchor._p.addnext(el)
    anchor=Paragraph(el,doc._body); set_text(anchor,text)
    anchor.paragraph_format.keep_together=True
    anchor.paragraph_format.keep_with_next=False
    anchor.paragraph_format.page_break_before=False

# Replace long runs of blank paragraphs used as manual page spacing.
for i in list(range(314,341))+list(range(346,373))+[383,384]:
    p=ps[i]
    if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:sectPr'):
        p._p.getparent().remove(p._p)
ps[341].paragraph_format.page_break_before=True
ps[373].paragraph_format.page_break_before=True

# Keep revised ERD and database headings with the following content.
for i in [272,282]: ps[i].paragraph_format.page_break_before=True
for style_name in ['Title','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[style_name]; style.font.color.rgb=RGBColor(0,0,0)
    for border in style.element.xpath('./w:pPr/w:pBdr'): border.getparent().remove(border)
doc.save(DEST)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
check=Document(DEST)
all_text='\n'.join(p.text for p in check.paragraphs)
assert 'sqlite' not in all_text.lower()
assert 'TAMBAH ERD SESSION' not in all_text
assert 'tidak menggunakan foreign key' not in all_text.lower()
assert len(check.tables)==22
assert all(f'[{i}]' in all_text for i in range(1,14))
(ASSETS/'changes.json').write_text(json.dumps({'source':str(SOURCE),'output':str(DEST),'source_sha256':source_hash,'changes':changes,'references':refs},ensure_ascii=False,indent=2),encoding='utf-8')
print('Revised',DEST)
print('Updated',len(changes),'paragraphs/captions, six database figures, and 13 bibliography entries.')
